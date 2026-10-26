from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
import sqlite3

from docx import Document
from pypdf import PdfWriter
import pytest

from resume_analyzer.analysis import ValidationError, text_stats, validate_text
from resume_analyzer.parsing import extract_resume, MAX_FILE_BYTES
from resume_analyzer.persistence import HistoryStore

SAMPLES = Path(__file__).resolve().parents[1] / "samples"


@pytest.mark.parametrize("extension", ["pdf", "docx"])
def test_fictional_extraction(extension):
    name = f"fictional_resume.{extension}"
    result = extract_resume(name, (SAMPLES / name).read_bytes())
    assert "Alex Example" in result.text
    assert "FICTIONAL" in result.text
    assert result.file_type == extension.upper()
    assert result.pages == (1 if extension == "pdf" else None)


@pytest.mark.parametrize("name,content", [("file.txt", b"hello"), ("file.pdf", b""),
    ("file.docx", b""), ("file.pdf", b"not pdf"), ("file.pdf", b"%PDF-broken"),
    ("file.docx", b"not docx"), ("file.pdf", b"x" * (MAX_FILE_BYTES + 1))],
    ids=["unsupported", "empty-pdf", "empty-docx", "bad-signature", "corrupt-pdf", "corrupt-docx", "oversized"])
def test_invalid_files(name, content):
    with pytest.raises(ValidationError):
        extract_resume(name, content)


@pytest.mark.parametrize("value", ["", " \n\t", "!!!", None, "x" * 100_001],
                         ids=["empty", "whitespace", "punctuation", "none", "too-long"])
def test_invalid_text(value):
    with pytest.raises(ValidationError):
        validate_text(value, "Job description")


def test_stats_and_unicode():
    stats = text_stats("Python python\nSQL café")
    assert stats.words == 4
    assert stats.unique_words == 3
    assert stats.nonempty_lines == 2
    assert stats.characters == 22


@pytest.mark.parametrize("extension", ["pdf", "docx"])
def test_no_readable_text(extension):
    output = BytesIO()
    if extension == "pdf":
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        writer.write(output)
    else:
        Document().save(output)
    with pytest.raises(ValidationError, match="No readable text"):
        extract_resume(f"blank.{extension}", output.getvalue())


def test_password_pdf():
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.encrypt("example-password")
    output = BytesIO()
    writer.write(output)
    with pytest.raises(ValidationError, match="Password-protected"):
        extract_resume("locked.pdf", output.getvalue())


def test_page_limit():
    writer = PdfWriter()
    for _ in range(51):
        writer.add_blank_page(width=100, height=100)
    output = BytesIO()
    writer.write(output)
    with pytest.raises(ValidationError, match="50 pages"):
        extract_resume("long.pdf", output.getvalue())


def test_docx_table():
    document = Document()
    document.add_paragraph("Fictional résumé")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Python SQL"
    output = BytesIO()
    document.save(output)
    assert "Python SQL" in extract_resume("resume.DOCX", output.getvalue()).text


def test_zip_disguised_as_docx():
    output = BytesIO()
    with ZipFile(output, "w") as archive:
        archive.writestr("other.txt", "hello")
    with pytest.raises(ValidationError, match="valid Word"):
        extract_resume("fake.docx", output.getvalue())


def test_history_privacy_and_roundtrip(tmp_path):
    path = tmp_path / "nested/history.sqlite3"
    store = HistoryStore(path)
    assert store.recent() == []
    resume = extract_resume("private-name.docx", (SAMPLES / "fictional_resume.docx").read_bytes())
    job = "private text'); DROP TABLE analyses; --"
    store.save(resume, text_stats(resume.text), text_stats(job))
    rows = store.recent()
    assert len(rows) == 1
    assert rows[0]["resume_words"] == text_stats(resume.text).words
    assert "Alex Example" not in str(rows)
    with sqlite3.connect(path) as connection:
        columns = {row[1] for row in connection.execute("PRAGMA table_info(analyses)")}
    assert not {"filename", "resume_text", "job_text"} & columns
    assert b"Alex Example" not in path.read_bytes()
    assert b"private text" not in path.read_bytes()
    store.clear()
    assert store.recent() == []
